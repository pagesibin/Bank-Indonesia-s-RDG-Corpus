"""Build out/labeling_sheet.xlsx — a blind, stratified hand-labelling sample for
validating the tone index.

- 200 Indonesian sections from out/corpus.csv, stratified by topic x era, weighted
  toward topics that carry a policy-stance signal (inflation, exchange rate, growth).
  Payment-system / digital-payment sections are left out (no rate-stance content).
- 20 of them are repeated later in the sheet (shuffled in) to measure the labeller's
  own consistency (intra-rater reliability).
- The Label sheet is BLIND: no date, code or decision is shown, so knowing the rate
  path can't bias the label. The hidden Key sheet maps each item back to its release.

    python make_labeling_sheet.py        # fixed seed -> same sample every run
"""
import random

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SEED = 20260930
OUT = "out/labeling_sheet.xlsx"
N_REPEAT = 20

# sections per topic (sums to 200); spread as evenly as possible over the 4 eras
QUOTA = {
    "inflation": 36, "exchange_rate": 32, "growth": 28, "global": 20,
    "decision": 16, "bop": 16, "credit": 14, "monetary_policy": 14,
    "banking_resilience": 6, "money_supply": 4, "macroprudential": 4, "other": 10,
}
ERAS = [(2005, 2011, "2005-11"), (2012, 2016, "2012-16"),
        (2017, 2020, "2017-20"), (2021, 2026, "2021-26")]

c = pd.read_csv("out/corpus.csv", dtype={"code": str})
c = c[(c.lang == "id") & c.n_words.between(20, 600)].copy()
c["year"] = c.date.str[:4].astype(int)
c["era"] = c.year.map(lambda y: next(l for a, b, l in ERAS if a <= y <= b))
c["full_text"] = (c.heading.fillna("").str.strip().str.rstrip(".:;") + ". " + c.text.fillna("").str.strip()).str.strip(". ")

rng = random.Random(SEED)
picked = []
for topic, n in QUOTA.items():
    pool = c[c.topic == topic]
    by_era = {e: pool[pool.era == e].index.tolist() for _, _, e in ERAS}
    for lst in by_era.values():
        rng.shuffle(lst)
    # round-robin across eras so each era gets an equal share where it has rows
    got = []
    while len(got) < n and any(by_era.values()):
        for e in list(by_era):
            if by_era[e] and len(got) < n:
                got.append(by_era[e].pop())
    picked += got

sample = c.loc[picked].copy()
sample = sample.sample(frac=1, random_state=SEED).reset_index(drop=True)
sample["item"] = [f"S{i + 1:03d}" for i in range(len(sample))]
sample["repeat_of"] = ""

# repeats: drawn from the first 150, inserted among the last 60 so they aren't adjacent
rep = sample.iloc[rng.sample(range(150), N_REPEAT)].copy()
rep["repeat_of"] = rep["item"]
rows = sample.to_dict("records")
for r in rep.to_dict("records"):
    rows.insert(rng.randint(len(rows) - 60, len(rows)), r)
for i, r in enumerate(rows):
    r["item"] = f"S{i + 1:03d}"
items = pd.DataFrame(rows)
items["repeat_of"] = items.apply(
    lambda r: items.loc[(items.code == r.code) & (items.section_order == r.section_order)
                        & (items.repeat_of == ""), "item"].iloc[0] if r.repeat_of else "", axis=1)

docs = pd.read_csv("out/documents.csv", dtype={"code": str})
docs = docs[docs.lang == "id"][["code", "direction", "bps", "rate_level"]]
items = items.drop(columns=["direction", "rate_level"], errors="ignore").merge(docs, on="code", how="left")

# ---------------------------------------------------------------------------- workbook
F = "Arial"
H_FILL = PatternFill("solid", start_color="1F3864")
IN_FILL = PatternFill("solid", start_color="FFF2CC")   # cells to fill in
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# --- Guide
g = wb.active
g.title = "Guide"
g.sheet_view.showGridLines = False
g.column_dimensions["A"].width = 26
g.column_dimensions["B"].width = 100
lines = [
    ("BI RDG tone — hand-labelling sheet", None, "title"),
    ("What to do", "Go to the 'Label' sheet. For every row, read the text and fill in the YELLOW cells: "
     "Label, Confidence, Forward-looking and (optionally) Notes. Nothing else needs editing.", None),
    ("The question", "Read as a bond trader would: does this paragraph point BI toward HIGHER rates / "
     "tighter policy (Hawkish), toward LOWER rates / easier policy (Dovish), or neither (Neutral)?", None),
    ("Blind by design", "Dates and decisions are hidden on purpose. Judge only the text in front of you; "
     "don't try to guess the year or look up the meeting.", None),
    ("", None, None),
    ("Labels", None, "head"),
    ("Hawkish", "Signals upward pressure on rates: inflation rising or above target, rupiah depreciating "
     "or under pressure, overheating demand/credit, capital outflows, BI 'vigilant' (waspada) about risks, "
     "stance described as ketat / pre-emptive / ahead of the curve.", None),
    ("Dovish", "Signals room for lower rates: inflation low or falling within target, rupiah stable or "
     "strengthening, growth weak or slowing, credit sluggish, stance described as longgar / akomodatif / "
     "mendorong pertumbuhan, explicit 'ruang penurunan' (room to cut).", None),
    ("Neutral", "Descriptive with no stance implication, balanced (risks both ways), or not about "
     "monetary conditions at all.", None),
    ("", None, None),
    ("Topic matters", None, "head"),
    ("Same word, different sign", "Read the direction word together with its topic. "
     "Inflasi / kredit / defisit meningkat → HAWKISH.  Pertumbuhan melambat / melemah → DOVISH; "
     "pertumbuhan kuat / tinggi → leans HAWKISH.  Rupiah melemah / terdepresiasi → HAWKISH; "
     "Rupiah menguat / stabil → leans DOVISH.", None),
    ("Negation", "'tidak meningkat', 'belum pulih' etc. reverse or qualify the direction word.", None),
    ("Decision paragraphs", "A paragraph announcing the rate decision: hike → Hawkish, cut → Dovish, "
     "hold → judge the reasoning that follows (a hold 'to anticipate inflation pressure' is Hawkish; "
     "a hold 'to support growth' is Dovish; a plain hold is Neutral).", None),
    ("", None, None),
    ("Other columns", None, "head"),
    ("Confidence", "1 = unsure / could argue either way · 2 = fairly sure · 3 = clear-cut.", None),
    ("Forward-looking", "Y if the paragraph is mainly about the future (diperkirakan, ke depan, akan, "
     "prospek, proyeksi); N if it mainly describes what already happened.", None),
    ("Notes", "Anything useful: ambiguous wording, mixed signals, a phrase that decided it for you.", None),
    ("", None, None),
    ("Example (not part of the sample)", None, "head"),
    ("Text", "Inflasi inti diperkirakan meningkat seiring dengan tekanan depresiasi nilai tukar Rupiah dan "
     "kenaikan ekspektasi inflasi.", None),
    ("Label / Conf. / Fwd", "Hawkish  ·  3  ·  Y", None),
    ("Why", "Core inflation expected to rise + rupiah depreciation pressure → both point to tighter policy.", None),
    ("", None, None),
    ("Tips", "About 220 rows ≈ 2 hours. Take breaks; your progress is shown below. Some items appear "
     "twice on purpose (to measure labelling consistency) — just label them again as you see them, "
     "without scrolling back.", None),
    ("", None, None),
    ("Progress", None, "head"),
]
r = 1
for a, b, kind in lines:
    ca = g.cell(row=r, column=1, value=a or None)
    if kind == "title":
        ca.font = Font(name=F, size=16, bold=True, color="1F3864")
    elif kind == "head":
        ca.font = Font(name=F, size=12, bold=True, color="1F3864")
    else:
        ca.font = Font(name=F, size=10, bold=True)
        ca.alignment = WRAP
    if b:
        cb = g.cell(row=r, column=2, value=b)
        cb.font = Font(name=F, size=10)
        cb.alignment = WRAP
        g.row_dimensions[r].height = max(15, 15 * (len(b) // 95 + 1))
    r += 1

n_items = len(items)
last = n_items + 1
prog = [
    ("Items in sample", f"=COUNTA(Label!A2:A{last})"),
    ("Labelled so far", f"=COUNTA(Label!D2:D{last})"),
    ("Remaining", f"=B{r}-B{r + 1}"),
    ("% done", f"=IF(B{r}=0,0,B{r + 1}/B{r})"),
    ("Hawkish", f'=COUNTIF(Label!D2:D{last},"Hawkish")'),
    ("Neutral", f'=COUNTIF(Label!D2:D{last},"Neutral")'),
    ("Dovish", f'=COUNTIF(Label!D2:D{last},"Dovish")'),
]
for i, (a, f) in enumerate(prog):
    g.cell(row=r + i, column=1, value=a).font = Font(name=F, size=10, bold=True)
    cell = g.cell(row=r + i, column=2, value=f)
    cell.font = Font(name=F, size=10)
    cell.alignment = Alignment(horizontal="left")
    if a == "% done":
        cell.number_format = "0%"

# --- Label
ws = wb.create_sheet("Label")
heads = ["Item", "Topic", "Text (heading + paragraph)", "Label", "Confidence (1-3)",
         "Forward-looking (Y/N)", "Notes", "Score"]
widths = [8, 16, 95, 12, 12, 13, 30, 8]
for j, (h, w) in enumerate(zip(heads, widths), 1):
    cell = ws.cell(row=1, column=j, value=h)
    cell.font = Font(name=F, size=10, bold=True, color="FFFFFF")
    cell.fill = H_FILL
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = BOX
    ws.column_dimensions[get_column_letter(j)].width = w
ws.row_dimensions[1].height = 30
ws.freeze_panes = "D2"
ws["D1"].comment = Comment("Pick from the list: Hawkish / Neutral / Dovish", "labeling sheet")
ws["H1"].comment = Comment("Filled automatically: Hawkish = +1, Neutral = 0, Dovish = -1", "labeling sheet")

topic_names = {
    "inflation": "Inflation", "exchange_rate": "Exchange rate", "growth": "Growth",
    "global": "Global", "decision": "Decision", "bop": "Balance of payments",
    "credit": "Credit", "monetary_policy": "Monetary policy",
    "banking_resilience": "Banking", "money_supply": "Money supply",
    "macroprudential": "Macroprudential", "other": "Other",
}
for i, row in enumerate(items.itertuples(), 2):
    vals = [row.item, topic_names.get(row.topic, row.topic), row.full_text]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=i, column=j, value=v)
        cell.font = Font(name=F, size=10)
        cell.alignment = WRAP
        cell.border = BOX
    for j in (4, 5, 6, 7):
        cell = ws.cell(row=i, column=j)
        cell.fill = IN_FILL
        cell.border = BOX
        cell.font = Font(name=F, size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="center" if j < 7 else "left")
    sc = ws.cell(row=i, column=8,
                 value=f'=IF(D{i}="Hawkish",1,IF(D{i}="Dovish",-1,IF(D{i}="Neutral",0,"")))')
    sc.font = Font(name=F, size=10)
    sc.alignment = Alignment(horizontal="center", vertical="top")
    sc.border = BOX
    ws.row_dimensions[i].height = min(400, 13 * (len(row.full_text) // 105 + 1) + 4)

for formula, col in [('"Hawkish,Neutral,Dovish"', "D"), ('"1,2,3"', "E"), ('"Y,N"', "F")]:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True,
                        showErrorMessage=True, errorTitle="Invalid entry",
                        error="Please pick a value from the drop-down list.")
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{last}")
ws.auto_filter.ref = f"A1:H{last}"

# --- Key (hidden)
k = wb.create_sheet("Key")
kcols = ["item", "repeat_of", "code", "lang", "date", "era", "section_order", "topic",
         "heading_key", "n_words", "direction", "bps", "rate_level"]
for j, h in enumerate(kcols, 1):
    cell = k.cell(row=1, column=j, value=h)
    cell.font = Font(name=F, size=10, bold=True)
for i, row in enumerate(items[kcols].itertuples(index=False), 2):
    for j, v in enumerate(row, 1):
        k.cell(row=i, column=j, value=None if pd.isna(v) else v).font = Font(name=F, size=10)
k.cell(row=last + 2, column=1,
       value="Hidden answer key: maps each item back to its release. 'repeat_of' marks the "
             f"{N_REPEAT} consistency repeats. Generated by make_labeling_sheet.py, seed {SEED}.").font = Font(name=F, size=9, italic=True)
k.sheet_state = "hidden"

wb.active = 1
wb.save(OUT)

print(f"{OUT}: {len(items)} rows ({len(items) - N_REPEAT} unique + {N_REPEAT} repeats)")
print(pd.crosstab(items[items.repeat_of == ''].topic, items[items.repeat_of == ''].era, margins=True))
